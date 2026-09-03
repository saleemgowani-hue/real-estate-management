"""
utils.py
Reusable, UI-agnostic helper functions: date/number formatting, validation,
Excel export (openpyxl), and PDF export (reportlab).
"""

import io
import datetime
import pandas as pd

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

from config import APP_NAME, APP_FOOTER


# ---------------------------------------------------------------------------
# Date / number helpers
# ---------------------------------------------------------------------------
def today_str():
    return datetime.date.today().strftime("%Y-%m-%d")


def now_str():
    return datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def fmt_date(value, date_format="DD-MM-YYYY"):
    """Safely format a date-like string/object; never raises on bad input."""
    if value in (None, "", "None"):
        return "-"
    try:
        if isinstance(value, (datetime.date, datetime.datetime)):
            d = value
        else:
            d = pd.to_datetime(str(value), errors="coerce")
            if pd.isna(d):
                return str(value)
        mapping = {
            "DD-MM-YYYY": "%d-%m-%Y",
            "MM-DD-YYYY": "%m-%d-%Y",
            "YYYY-MM-DD": "%Y-%m-%d",
        }
        return d.strftime(mapping.get(date_format, "%d-%m-%Y"))
    except Exception:
        return str(value)


def fmt_currency(value, symbol="₹"):
    try:
        value = float(value or 0)
        return f"{symbol}{value:,.2f}"
    except (TypeError, ValueError):
        return f"{symbol}0.00"


def safe_float(value, default=0.0):
    try:
        if value in (None, "", "None"):
            return default
        return float(value)
    except (TypeError, ValueError):
        return default


def safe_int(value, default=0):
    try:
        if value in (None, "", "None"):
            return default
        return int(value)
    except (TypeError, ValueError):
        return default


def date_range_for_filter(filter_name):
    """Returns (start_date, end_date) as date objects for a named report filter."""
    today = datetime.date.today()
    if filter_name == "Today":
        return today, today
    if filter_name == "Yesterday":
        y = today - datetime.timedelta(days=1)
        return y, y
    if filter_name == "This Week":
        start = today - datetime.timedelta(days=today.weekday())
        return start, today
    if filter_name == "This Month":
        return today.replace(day=1), today
    if filter_name == "Last Month":
        first_this_month = today.replace(day=1)
        last_month_end = first_this_month - datetime.timedelta(days=1)
        return last_month_end.replace(day=1), last_month_end
    if filter_name == "This Year":
        return today.replace(month=1, day=1), today
    return None, None  # Custom / All Time — caller supplies dates


# ---------------------------------------------------------------------------
# Excel export
# ---------------------------------------------------------------------------
def dataframe_to_excel_bytes(df: pd.DataFrame, sheet_name="Report", title=None) -> bytes:
    """Returns raw xlsx bytes for a styled export of the given DataFrame."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name[:31] if sheet_name else "Report"

    start_row = 1
    if title:
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(1, len(df.columns)))
        cell = ws.cell(row=1, column=1, value=title)
        cell.font = Font(bold=True, size=14, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4F46E5")
        cell.alignment = Alignment(horizontal="center")
        start_row = 3

    if df is None or df.empty:
        ws.cell(row=start_row, column=1, value="No data available.")
    else:
        header_row = start_row
        for c_idx, col in enumerate(df.columns, start=1):
            cell = ws.cell(row=header_row, column=c_idx, value=str(col))
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="0EA5E9")
            cell.alignment = Alignment(horizontal="center")

        for r_idx, row in enumerate(df.itertuples(index=False), start=header_row + 1):
            for c_idx, value in enumerate(row, start=1):
                ws.cell(row=r_idx, column=c_idx, value=value)

        for c_idx, col in enumerate(df.columns, start=1):
            max_len = max([len(str(col))] + [len(str(v)) for v in df[col].astype(str).values[:200]])
            ws.column_dimensions[get_column_letter(c_idx)].width = min(max(12, max_len + 2), 45)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# PDF export
# ---------------------------------------------------------------------------
def dataframe_to_pdf_bytes(df: pd.DataFrame, title="Report", company_name=APP_NAME) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=landscape(A4),
        leftMargin=1.2 * cm, rightMargin=1.2 * cm, topMargin=1.2 * cm, bottomMargin=1.2 * cm,
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("TitleStyle", parent=styles["Heading1"], textColor=colors.HexColor("#4F46E5"))
    sub_style = ParagraphStyle("SubStyle", parent=styles["Normal"], textColor=colors.grey)

    elements = [
        Paragraph(company_name, title_style),
        Paragraph(title, styles["Heading2"]),
        Paragraph(f"Generated on {datetime.datetime.now().strftime('%d-%m-%Y %H:%M')}", sub_style),
        Spacer(1, 0.5 * cm),
    ]

    if df is None or df.empty:
        elements.append(Paragraph("No data available for the selected filters.", styles["Normal"]))
    else:
        data = [list(df.columns)] + df.astype(str).values.tolist()
        table = Table(data, repeatRows=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4F46E5")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 7.5),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CBD5E1")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F5F9")]),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("ALIGN", (0, 0), (-1, -1), "LEFT"),
        ]))
        elements.append(table)

    elements.append(Spacer(1, 0.8 * cm))
    elements.append(Paragraph(APP_FOOTER, sub_style))

    doc.build(elements)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Validation helpers
# ---------------------------------------------------------------------------
def is_blank(value):
    return value is None or str(value).strip() == ""
